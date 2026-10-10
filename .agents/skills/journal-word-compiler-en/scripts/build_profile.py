from __future__ import annotations

import argparse
import html
import re
import shutil
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from common import (
    dump_json,
    engine_root_default,
    ensure_engine,
    file_record,
    load_json,
    slugify,
    validate_ooxml_input,
)


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)


def extract_guide_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(value for value in parts if value.strip())
    if suffix in {".html", ".htm"}:
        parser = _TextHTMLParser()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        return "\n".join(html.unescape(value) for value in parser.parts)
    if suffix == ".pdf":
        try:
            import fitz

            with fitz.open(path) as pdf:
                return "\n".join(page.get_text("text") for page in pdf)
        except Exception:
            completed = subprocess.run(
                ["pdftotext", "-layout", str(path), "-"],
                text=True,
                encoding="utf-8",
                errors="replace",
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                check=False,
            )
            if completed.returncode == 0:
                return completed.stdout
            raise RuntimeError(f"Unable to extract guide PDF: {completed.stderr}")
    raise ValueError(f"Unsupported guide type: {suffix}")


def extract_rule_candidates(text: str) -> list[dict[str, Any]]:
    patterns = [
        ("abstract_words_max", r"abstract[^\n]{0,100}?(\d{2,4})\s*words?"),
        ("keywords_range", r"(\d{1,2})\s*(?:to|[-–—])\s*(\d{1,2})\s*(?:keywords?|key words?|terms?)"),
        ("keywords_max", r"(?:maximum|max|up to|no more than)[^\n]{0,40}?(\d{1,2})\s*(?:keywords?|key words?|terms?)"),
        ("font_points", r"(?:font|typeface|text)[^\n]{0,100}?(\d{1,2}(?:\.5)?)\s*(?:pt|point)"),
        ("line_spacing", r"(?:line spacing|spaced)[^\n]{0,80}?(single|double|1\.5)"),
        ("margin", r"margins?[^\n]{0,80}?(\d+(?:\.\d+)?)\s*(cm|mm|inches?|inch)"),
    ]
    candidates: list[dict[str, Any]] = []
    for name, pattern in patterns:
        for match in re.finditer(pattern, text, flags=re.I):
            start = max(0, text.rfind("\n", 0, match.start()) + 1)
            end_pos = text.find("\n", match.end())
            end = len(text) if end_pos < 0 else end_pos
            candidates.append(
                {
                    "kind": name,
                    "captures": list(match.groups()),
                    "source_excerpt": re.sub(r"\s+", " ", text[start:end]).strip()[:400],
                    "status": "candidate_only_not_applied",
                }
            )
    return candidates


def deep_merge(base: dict[str, Any], patch: dict[str, Any]) -> dict[str, Any]:
    merged = dict(base)
    for key, value in patch.items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key] = deep_merge(merged[key], value)
        else:
            merged[key] = value
    return merged


def extract_template_spec(
    template: Path,
    template_rolemap_path: Path,
    engine_root: Path,
    guide_rules_path: Path | None = None,
) -> dict[str, Any]:
    ensure_engine(engine_root)
    sys.path.insert(0, str(engine_root))
    try:
        from core.rules_from_template import extract_rules_from_template
        from core.schema import validate_spec

        rolemap_payload = load_json(template_rolemap_path)
        rolemap = {int(key): value for key, value in rolemap_payload.items()}
        spec = extract_rules_from_template(str(template), rolemap)
        if guide_rules_path:
            rules = load_json(guide_rules_path)
            if not isinstance(rules, dict):
                raise ValueError("guide-rules JSON must be an object containing FormatSpec overrides")
            spec = deep_merge(spec, rules)
        spec.setdefault("profile", "english_academic")
        spec.setdefault("locale", "en-US")
        spec["cleanup"] = {"mode": "preserve_emphasis"}
        validate_spec(spec)
        return spec
    finally:
        try:
            sys.path.remove(str(engine_root))
        except ValueError:
            pass


def build_profile(
    journal: str,
    template: Path,
    guide: Path,
    template_rolemap_path: Path,
    profile_store: Path,
    engine_root: Path,
    guide_rules_path: Path | None = None,
) -> Path:
    validate_ooxml_input(template, {".docx", ".dotx"})
    if guide.suffix.lower() not in {".pdf", ".docx", ".html", ".htm"}:
        raise ValueError("Guide must be PDF, DOCX, or saved HTML")
    if not guide.is_file():
        raise FileNotFoundError(guide)

    template_record = file_record(template)
    guide_record = file_record(guide)
    rules_record = file_record(guide_rules_path) if guide_rules_path else None
    rules_suffix = f"-{rules_record['sha256'][:8]}" if rules_record else "-template-only"
    version_id = (
        f"{template_record['sha256'][:8]}-{guide_record['sha256'][:8]}{rules_suffix}"
    )
    profile_dir = profile_store / slugify(journal) / version_id
    profile_dir.mkdir(parents=True, exist_ok=True)

    template_copy = profile_dir / f"template{template.suffix.lower()}"
    guide_copy = profile_dir / f"submission-guide{guide.suffix.lower()}"
    if not template_copy.exists():
        shutil.copy2(template, template_copy)
    if not guide_copy.exists():
        shutil.copy2(guide, guide_copy)
    shutil.copy2(template_rolemap_path, profile_dir / "template-rolemap.json")
    if guide_rules_path:
        shutil.copy2(guide_rules_path, profile_dir / "guide-rules.json")

    guide_text = extract_guide_text(guide)
    (profile_dir / "guide-extracted.txt").write_text(guide_text, encoding="utf-8")
    dump_json(profile_dir / "rule-candidates.json", extract_rule_candidates(guide_text))

    spec = extract_template_spec(
        template, template_rolemap_path, engine_root, guide_rules_path
    )
    dump_json(profile_dir / "profile.json", spec)
    dump_json(
        profile_dir / "source-manifest.json",
        {
            "journal": journal,
            "version_id": version_id,
            "template": template_record,
            "submission_guide": guide_record,
            "guide_rules": rules_record,
            "rule_precedence": [
                "explicit reviewed guide-rules JSON",
                "official Word template",
                "unapplied extracted rule candidates",
            ],
        },
    )
    exceptions = [
        "# Exceptions and review notes",
        "",
        "- Automatically extracted guide candidates are evidence only and are not applied until reviewed.",
        "- Conflicts between an explicit guide rule and the template must be resolved before submission.",
        "- Citation wording and reference metadata are preserved, not rewritten.",
    ]
    (profile_dir / "exceptions.md").write_text("\n".join(exceptions) + "\n", encoding="utf-8")
    return profile_dir


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a versioned English journal profile.")
    parser.add_argument("--journal", required=True)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--guide", required=True, type=Path)
    parser.add_argument("--template-rolemap", required=True, type=Path)
    parser.add_argument("--profile-store", required=True, type=Path)
    parser.add_argument("--guide-rules", type=Path)
    parser.add_argument("--engine-root", type=Path, default=engine_root_default())
    args = parser.parse_args()
    profile = build_profile(
        args.journal,
        args.template,
        args.guide,
        args.template_rolemap,
        args.profile_store,
        args.engine_root,
        args.guide_rules,
    )
    print(profile)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
